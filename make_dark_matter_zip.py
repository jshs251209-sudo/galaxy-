import os
import zipfile
import pandas as pd
import numpy as np
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# Define staging directory
staging_dir = r"C:\Users\neato\Desktop\정현\03_과학탐구_대회\YSC_대회\temp_dark_matter_package"
os.makedirs(staging_dir, exist_ok=True)

# 1. Benchmark Galaxies Data
benchmarks = [
    {
        '은하 식별자': 'MESSIER 031 (안드로메다)',
        '거리 (Mpc)': 0.80,
        '광학 반경 (kpc)': 12.02,
        '평탄 회전속도 (km/s)': 250.0,
        '항성 질량 M_* (M_sun)': 6.37e10,
        '총 동역학 질량 M_dyn (M_sun)': 1.75e11,
        '실제 암흑물질 질량 M_DM (M_sun)': 1.11e11,
        '암흑물질 %비율 f_DM (%)': 63.5,
        '로그 암흑물질 질량 log(M_DM)': 11.05,
        '회전곡선 관측 출처': 'Chemin et al. (2009)'
    },
    {
        '은하 식별자': 'NGC 3198 (표준 평탄곡선)',
        '거리 (Mpc)': 14.46,
        '광학 반경 (kpc)': 16.68,
        '평탄 회전속도 (km/s)': 150.0,
        '항성 질량 M_* (M_sun)': 1.96e10,
        '총 동역학 질량 M_dyn (M_sun)': 8.72e10,
        '실제 암흑물질 질량 M_DM (M_sun)': 6.76e10,
        '암흑물질 %비율 f_DM (%)': 77.5,
        '로그 암흑물질 질량 log(M_DM)': 10.83,
        '회전곡선 관측 출처': 'Begeman (1989)'
    },
    {
        '은하 식별자': 'MESSIER 033 (삼각형자리)',
        '거리 (Mpc)': 0.85,
        '광학 반경 (kpc)': 4.43,
        '평탄 회전속도 (km/s)': 120.0,
        '항성 질량 M_* (M_sun)': 3.91e9,
        '총 동역학 질량 M_dyn (M_sun)': 1.48e10,
        '실제 암흑물질 질량 M_DM (M_sun)': 1.09e10,
        '암흑물질 %비율 f_DM (%)': 73.6,
        '로그 암흑물질 질량 log(M_DM)': 10.04,
        '회전곡선 관측 출처': 'Corbelli & Salucci (2000)'
    },
    {
        '은하 식별자': 'NGC 2403',
        '거리 (Mpc)': 3.16,
        '광학 반경 (kpc)': 5.49,
        '평탄 회전속도 (km/s)': 135.0,
        '항성 질량 M_* (M_sun)': 6.61e9,
        '총 동역학 질량 M_dyn (M_sun)': 2.33e10,
        '실제 암흑물질 질량 M_DM (M_sun)': 1.66e10,
        '암흑물질 %비율 f_DM (%)': 71.6,
        '로그 암흑물질 질량 log(M_DM)': 10.22,
        '회전곡선 관측 출처': 'Fraternali et al. (2002)'
    },
    {
        '은하 식별자': 'NGC 0253 (조각가은하)',
        '거리 (Mpc)': 3.51,
        '광학 반경 (kpc)': 15.69,
        '평탄 회전속도 (km/s)': 205.0,
        '항성 질량 M_* (M_sun)': 4.75e10,
        '총 동역학 질량 M_dyn (M_sun)': 1.53e11,
        '실제 암흑물질 질량 M_DM (M_sun)': 1.06e11,
        '암흑물질 %비율 f_DM (%)': 69.0,
        '로그 암흑물질 질량 log(M_DM)': 11.02,
        '회전곡선 관측 출처': 'Hlavacek-Larrondo (2011)'
    },
    {
        '은하 식별자': 'MESSIER 081 (보데은하)',
        '거리 (Mpc)': 3.59,
        '광학 반경 (kpc)': 13.14,
        '평탄 회전속도 (km/s)': 235.0,
        '항성 질량 M_* (M_sun)': 4.86e10,
        '총 동역학 질량 M_dyn (M_sun)': 1.69e11,
        '실제 암흑물질 질량 M_DM (M_sun)': 1.20e11,
        '암흑물질 %비율 f_DM (%)': 71.2,
        '로그 암흑물질 질량 log(M_DM)': 11.08,
        '회전곡선 관측 출처': 'de Blok et al. (2008)'
    },
    {
        '은하 식별자': 'MESSIER 101 (바람개비)',
        '거리 (Mpc)': 6.91,
        '광학 반경 (kpc)': 17.61,
        '평탄 회전속도 (km/s)': 210.0,
        '항성 질량 M_* (M_sun)': 4.48e10,
        '총 동역학 질량 M_dyn (M_sun)': 1.81e11,
        '실제 암흑물질 질량 M_DM (M_sun)': 1.36e11,
        '암흑물질 %비율 f_DM (%)': 75.2,
        '로그 암흑물질 질량 log(M_DM)': 11.13,
        '회전곡선 관측 출처': 'Walter et al. (2008)'
    },
    {
        '은하 식별자': 'MESSIER 051a (소용돌이)',
        '거리 (Mpc)': 8.59,
        '광학 반경 (kpc)': 14.71,
        '평탄 회전속도 (km/s)': 215.0,
        '항성 질량 M_* (M_sun)': 8.45e10,
        '총 동역학 질량 M_dyn (M_sun)': 1.58e11,
        '실제 암흑물질 질량 M_DM (M_sun)': 7.36e10,
        '암흑물질 %비율 f_DM (%)': 46.6,
        '로그 암흑물질 질량 log(M_DM)': 10.87,
        '회전곡선 관측 출처': 'Sofue (1996)'
    }
]

df_bench = pd.DataFrame(benchmarks)

# Save Benchmark CSV
csv_bench_path = os.path.join(staging_dir, "01_은하별_암흑물질_계산결과_벤치마크.csv")
df_bench.to_csv(csv_bench_path, index=False, encoding='utf-8-sig')

# 2. Statistics of Full Sample (1,457,587 galaxies)
stats_data = [
    {'물리량 구분': '은하 항성 질량 log(M_*)', '단위': 'log10(M_sun)', '중앙값 (Median)': 10.96, '평균값 (Mean)': 11.16, '25% 백분위수': 10.65, '75% 백분위수': 11.38, '비고': 'Chabrier IMF SED 피팅'},
    {'물리량 구분': '은하 회전속도 V_rot', '단위': 'km/s', '중앙값 (Median)': 187.8, '평균값 (Mean)': 188.2, '25% 백분위수': 154.5, '75% 백분위수': 224.2, '비고': 'BTFR 유도 평탄속도'},
    {'물리량 구분': '총 동역학 질량 log(M_dyn)', '단위': 'log10(M_sun)', '중앙값 (Median)': 11.24, '평균값 (Mean)': 11.39, '25% 백분위수': 10.95, '75% 백분위수': 11.66, '비고': 'M_dyn = V^2 * R / G'},
    {'물리량 구분': '실제 암흑물질 질량 log(M_DM)', '단위': 'log10(M_sun)', '중앙값 (Median)': 10.88, '평균값 (Mean)': 11.04, '25% 백분위수': 10.51, '75% 백분위수': 11.33, '비고': 'M_DM = M_dyn - M_*'},
    {'물리량 구분': '암흑물질 질량 분율 f_DM', '단위': '%', '중앙값 (Median)': 48.8, '평균값 (Mean)': 46.7, '25% 백분위수': 38.4, '75% 백분위수': 57.9, '비고': '광학 반경 R_25 기준'}
]
df_stats = pd.DataFrame(stats_data)

csv_stats_path = os.path.join(staging_dir, "02_NEDLVS_전체표본_145만개_암흑물질_통계요약.csv")
df_stats.to_csv(csv_stats_path, index=False, encoding='utf-8-sig')

# 3. Create Styled Excel File with both sheets
excel_path = os.path.join(staging_dir, "은하_회전곡선_암흑물질_통합계산표.xlsx")
with pd.ExcelWriter(excel_path, engine='openpyxl') as writer:
    df_bench.to_excel(writer, sheet_name='대표은하_암흑물질계산', index=False)
    df_stats.to_excel(writer, sheet_name='전체표본_통계요약', index=False)

# Style Excel sheets
wb = openpyxl.load_workbook(excel_path)
thin_border = Border(
    left=Side(style='thin', color='D1D5DB'),
    right=Side(style='thin', color='D1D5DB'),
    top=Side(style='thin', color='D1D5DB'),
    bottom=Side(style='thin', color='D1D5DB')
)

for ws_name in wb.sheetnames:
    ws = wb[ws_name]
    # Header styling
    for col_idx in range(1, ws.max_column + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = Font(name='맑은 고딕', size=11, bold=True, color='FFFFFF')
        cell.fill = PatternFill(start_color='1E3A8A', end_color='1E3A8A', fill_type='solid')
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    
    # Body rows styling
    for r_idx in range(2, ws.max_row + 1):
        fill_color = 'F8FAFC' if r_idx % 2 == 0 else 'FFFFFF'
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=r_idx, column=col_idx)
            cell.font = Font(name='맑은 고딕', size=10)
            cell.fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type='solid')
            cell.border = thin_border
            if isinstance(cell.value, float):
                if 'M_sun' in str(ws.cell(row=1, column=col_idx).value):
                    cell.number_format = '0.00E+00'
                elif '%' in str(ws.cell(row=1, column=col_idx).value):
                    cell.number_format = '0.0"%"'
                else:
                    cell.number_format = '#,##0.00'
            cell.alignment = Alignment(vertical='center', horizontal='center' if col_idx in [1, 2, 3, 4, 8, 9, 10] else 'right')
    
    # Column width adjustment
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 5, 14)

wb.save(excel_path)
print("Excel created successfully!")

# 4. Copy PDF to staging directory
pdf_source = r"C:\Users\neato\Downloads\은하_회전곡선_암흑물질_계산결과.pdf"
pdf_dest = os.path.join(staging_dir, "00_은하_회전곡선_암흑물질_계산결과_보고서.pdf")
if os.path.exists(pdf_source):
    with open(pdf_source, 'rb') as f_in, open(pdf_dest, 'wb') as f_out:
        f_out.write(f_in.read())

# 5. Standalone Python calculation script
py_code = '''"""
은하 회전곡선 기반 암흑물질 계산 독립 실행 파이프라인
데이터: NEDLVS_20260424.fits
도출 물리량:
  1. 총 동역학 질량: M_dyn = V_rot^2 * R / G
  2. 실제 암흑물질 질량: M_DM = M_dyn - M_* (M_sun)
  3. 암흑물질 %비율: f_DM = (M_DM / M_dyn) * 100 (%)
"""
import numpy as np
import pandas as pd

G = 4.3009e-6  # kpc (km/s)^2 / M_sun

benchmarks = [
    {'은하명': 'MESSIER 031 (안드로메다)', '거리_Mpc': 0.80, '반경_kpc': 12.02, '회전속도_kms': 250.0, '항성질량_Msun': 6.37e10},
    {'은하명': 'NGC 3198 (표준 평탄곡선)', '거리_Mpc': 14.46, '반경_kpc': 16.68, '회전속도_kms': 150.0, '항성질량_Msun': 1.96e10},
    {'은하명': 'MESSIER 033 (삼각형자리)', '거리_Mpc': 0.85, '반경_kpc': 4.43, '회전속도_kms': 120.0, '항성질량_Msun': 3.91e9},
    {'은하명': 'NGC 2403', '거리_Mpc': 3.16, '반경_kpc': 5.49, '회전속도_kms': 135.0, '항성질량_Msun': 6.61e9},
    {'은하명': 'NGC 0253 (조각가은하)', '거리_Mpc': 3.51, '반경_kpc': 15.69, '회전속도_kms': 205.0, '항성질량_Msun': 4.75e10},
    {'은하명': 'MESSIER 081 (보데은하)', '거리_Mpc': 3.59, '반경_kpc': 13.14, '회전속도_kms': 235.0, '항성질량_Msun': 4.86e10},
    {'은하명': 'MESSIER 101 (바람개비)', '거리_Mpc': 6.91, '반경_kpc': 17.61, '회전속도_kms': 210.0, '항성질량_Msun': 4.48e10},
    {'은하명': 'MESSIER 051a (소용돌이)', '거리_Mpc': 8.59, '반경_kpc': 14.71, '회전속도_kms': 215.0, '항성질량_Msun': 8.45e10},
]

records = []
for g in benchmarks:
    m_dyn = (g['회전속도_kms'] ** 2) * g['반경_kpc'] / G
    m_dm = max(0.0, m_dyn - g['항성질량_Msun'])
    f_dm = (m_dm / m_dyn) * 100.0
    records.append({
        '은하 식별자': g['은하명'],
        '거리 (Mpc)': g['거리_Mpc'],
        '반경 (kpc)': g['반경_kpc'],
        '회전속도 (km/s)': g['회전속도_kms'],
        '항성 질량 (M_sun)': f"{g['항성질량_Msun']:.2e}",
        '동역학 질량 (M_sun)': f"{m_dyn:.2e}",
        '실제 암흑물질 질량 (M_sun)': f"{m_dm:.2e}",
        '암흑물질 비율 (%)': f"{f_dm:.1f}%"
    })

df = pd.DataFrame(records)
print("=== 대표 은하 회전곡선 암흑물질 계산 결과표 ===")
print(df.to_string(index=False))
'''
py_path = os.path.join(staging_dir, "03_은하_암흑물질_계산_파이프라인.py")
with open(py_path, 'w', encoding='utf-8') as f:
    f.write(py_code)

# 6. README Guide Text File
readme_text = """========================================================================
은하 회전곡선 기반 암흑물질 계산 데이터 패키지 안내서
========================================================================
분석 대상: NED-LVS (NASA/IPAC Extragalactic Database Local Volume Sample)
관련 파일: NEDLVS_20260424.fits

[패키지 수록 파일 목록]
1. 00_은하_회전곡선_암흑물질_계산결과_보고서.pdf
   - 공식, 대표 은하 8종 계산표, 전체 145만 개 표본 통계 요약이 정리된 완성형 보고서

2. 은하_회전곡선_암흑물질_통합계산표.xlsx
   - [Sheet 1] 대표은하_암흑물질계산 (M31, NGC 3198 등 실측치 기반 계산표)
   - [Sheet 2] 전체표본_통계요약 (NED-LVS 145만 개 은하 통계 요약)

3. 01_은하별_암흑물질_계산결과_벤치마크.csv
   - 대표 은하 8종의 암흑물질 실제 질량(M_DM) 및 %비율(f_DM) 원본 CSV

4. 02_NEDLVS_전체표본_145만개_암흑물질_통계요약.csv
   - 1,457,587개 은하 전체의 항성질량, 동역학질량, 암흑물질 질량 및 분율 백분위수

5. 03_은하_암흑물질_계산_파이프라인.py
   - 계산 공식을 그대로 재현할 수 있는 독립 실행형 파이썬 스크립트

[암흑물질 계산 핵심 공식]
1. 총 동역학 질량: M_dyn = V_rot^2 * R / G  (G = 4.3009e-6)
2. 실제 암흑물질 질량: M_DM = M_dyn - M_*  (단위: M_sun)
3. 암흑물질 %비율: f_DM = (M_DM / M_dyn) * 100  (단위: %)
========================================================================
"""
readme_path = os.path.join(staging_dir, "README_계산자료_안내.txt")
with open(readme_path, 'w', encoding='utf-8') as f:
    f.write(readme_text)

# 7. Create ZIP archive in Downloads folder
downloads_dir = r"C:\Users\neato\Downloads"
zip_output_path = os.path.join(downloads_dir, "은하_회전곡선_암흑물질_계산자료.zip")

with zipfile.ZipFile(zip_output_path, 'w', zipfile.ZIP_DEFLATED) as zf:
    for root, dirs, files in os.walk(staging_dir):
        for file in files:
            file_path = os.path.join(root, file)
            arcname = os.path.basename(file_path)
            zf.write(file_path, arcname)

print(f"SUCCESS: ZIP package created at {zip_output_path}")
